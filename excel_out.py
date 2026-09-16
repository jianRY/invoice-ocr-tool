#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Excel 导出：票据汇总 + 明细 双表，全静态值。"""
import os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=10)
TOTAL_FILL = PatternFill("solid", fgColor="D9E2F3")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")
MONEY = "#,##0.00"

SUMMARY_HEADERS = ["文件名", "发票号码", "开票日期", "购买方", "销售方",
                   "项目数", "项目金额合计", "价税合计", "备注"]
DETAIL_HEADERS = ["文件名", "发票号码", "开票日期", "购买方", "销售方",
                  "序号", "项目名称", "数量", "项目金额", "价税合计(票)"]


def _style(ws, row, ncols, header=False, total=False):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.border = BORDER
        cell.font = HEADER_FONT if header else BODY_FONT
        if header:
            cell.fill = HEADER_FILL
            cell.alignment = CENTER
        elif total:
            cell.fill = TOTAL_FILL
        elif c >= 6:
            cell.alignment = RIGHT


def _autofit(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def export_excel(records, out_path):
    wb = Workbook()
    ws_s = wb.active
    ws_s.title = "票据汇总"
    ws_m = wb.create_sheet("明细")

    ws_m.append(DETAIL_HEADERS)
    _style(ws_m, 1, len(DETAIL_HEADERS), header=True)
    r = 2
    for rec in records:
        for idx, it in enumerate(rec["items"], 1):
            ws_m.append([rec["file"], rec["invoice_no"], rec["date"],
                         rec["buyer"], rec["seller"], idx, it["name"],
                         it.get("qty", ""), it["amount"], rec["total"]])
            _style(ws_m, r, len(DETAIL_HEADERS))
            ws_m.cell(row=r, column=9).number_format = MONEY
            ws_m.cell(row=r, column=10).number_format = MONEY
            r += 1
    _autofit(ws_m, {"A": 40, "B": 24, "C": 12, "D": 14, "E": 26, "F": 6,
                    "G": 42, "H": 10, "I": 12, "J": 13})
    ws_m.freeze_panes = "A2"
    ws_m.auto_filter.ref = f"A1:J{max(r - 1, 1)}"

    ws_s.append(SUMMARY_HEADERS)
    _style(ws_s, 1, len(SUMMARY_HEADERS), header=True)
    r = 2
    for rec in records:
        item_sum = sum(it["amount"] for it in rec["items"])
        ws_s.append([rec["file"], rec["invoice_no"], rec["date"],
                     rec["buyer"], rec["seller"], len(rec["items"]),
                     item_sum, rec["total"], rec.get("reason", "")])
        _style(ws_s, r, len(SUMMARY_HEADERS))
        ws_s.cell(row=r, column=7).number_format = MONEY
        ws_s.cell(row=r, column=8).number_format = MONEY
        r += 1
    ws_s.append(["合计", "", "", "", "",
                 sum(len(x["items"]) for x in records),
                 round(sum(sum(it["amount"] for it in x["items"]) for x in records), 2),
                 round(sum(x["total"] or 0 for x in records), 2), ""])
    _style(ws_s, r, len(SUMMARY_HEADERS), total=True)
    ws_s.cell(row=r, column=7).number_format = MONEY
    ws_s.cell(row=r, column=8).number_format = MONEY
    _autofit(ws_s, {"A": 40, "B": 24, "C": 12, "D": 14, "E": 26,
                    "F": 8, "G": 13, "H": 13, "I": 28})
    ws_s.freeze_panes = "A2"
    ws_s.auto_filter.ref = f"A1:I{max(r - 1, 1)}"

    wb.save(out_path)

    # 读回复核
    wb2 = load_workbook(out_path)
    s, m = wb2["票据汇总"], wb2["明细"]
    n_s, n_m = s.max_row - 2, m.max_row - 1
    return {"tickets": n_s, "details": n_m,
            "total": round(sum(x["total"] or 0 for x in records), 2),
            "size": os.path.getsize(out_path)}
